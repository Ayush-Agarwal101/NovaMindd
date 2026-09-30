"""
NovaMindd — DOCX Renderer

Deterministically renders a DocumentIR into a .docx file using python-docx.
"""
from __future__ import annotations

import io

from core.artifacts.ir import BlockType, DocumentIR
from core.logging import get_logger

logger = get_logger(__name__)


class DOCXRenderer:
    """Render DocumentIR → bytes (DOCX format)."""

    def render(self, doc_ir: DocumentIR) -> bytes:
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            raise RuntimeError("python-docx is required (pip install python-docx)")

        doc = Document()

        # Title
        title = doc.add_heading(doc_ir.title, level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Author metadata
        meta = doc.core_properties
        meta.author = doc_ir.author

        for block in doc_ir.blocks:
            if block.block_type == BlockType.HEADING1:
                doc.add_heading(block.text, level=1)
            elif block.block_type == BlockType.HEADING2:
                doc.add_heading(block.text, level=2)
            elif block.block_type == BlockType.HEADING3:
                doc.add_heading(block.text, level=3)
            elif block.block_type == BlockType.PARAGRAPH:
                doc.add_paragraph(block.text)
            elif block.block_type == BlockType.LIST:
                for item in block.items:
                    doc.add_paragraph(item, style="List Bullet")
            elif block.block_type == BlockType.CODE:
                p = doc.add_paragraph()
                run = p.add_run(block.text)
                run.font.name = "Courier New"
                run.font.size = Pt(9)
            elif block.block_type == BlockType.TABLE:
                if block.rows:
                    tbl = doc.add_table(
                        rows=len(block.rows),
                        cols=max(len(r) for r in block.rows),
                    )
                    tbl.style = "Table Grid"
                    for row_idx, row in enumerate(block.rows):
                        for col_idx, cell in enumerate(row):
                            c = tbl.cell(row_idx, col_idx)
                            c.text = cell.text
                            if cell.bold:
                                for para in c.paragraphs:
                                    for run in para.runs:
                                        run.bold = True
            elif block.block_type == BlockType.PAGE_BREAK:
                doc.add_page_break()

        # Evidence references appendix
        if doc_ir.evidence_refs:
            doc.add_heading("Evidence References", level=2)
            for ref in doc_ir.evidence_refs:
                doc.add_paragraph(
                    f"• {ref.get('document_id', '?')} — "
                    f"Page {ref.get('page', '?')} — "
                    f"Score {ref.get('score', 0):.3f}",
                    style="List Bullet",
                )

        buf = io.BytesIO()
        doc.save(buf)
        logger.debug("docx_rendered", title=doc_ir.title, blocks=len(doc_ir.blocks))
        return buf.getvalue()
