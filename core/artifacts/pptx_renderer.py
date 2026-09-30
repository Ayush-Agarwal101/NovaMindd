"""
NovaMindd — PPTX Renderer

Renders a DocumentIR into a .pptx file using python-pptx.
Headings become slide titles; paragraphs and lists become content.
"""
from __future__ import annotations

import io

from core.artifacts.ir import BlockType, DocumentIR
from core.logging import get_logger

logger = get_logger(__name__)


class PPTXRenderer:
    """Render DocumentIR → bytes (PPTX format)."""

    def render(self, doc_ir: DocumentIR) -> bytes:
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.enum.text import PP_ALIGN
        except ImportError:
            raise RuntimeError("python-pptx is required (pip install python-pptx)")

        prs = Presentation()
        blank_layout = prs.slide_layouts[6]    # Blank
        title_layout = prs.slide_layouts[0]    # Title Slide
        content_layout = prs.slide_layouts[1]  # Title and Content

        # Title slide
        title_slide = prs.slides.add_slide(title_layout)
        title_slide.shapes.title.text = doc_ir.title
        if len(title_slide.placeholders) > 1:
            title_slide.placeholders[1].text = f"Author: {doc_ir.author}"

        current_slide = None
        current_tf = None

        for block in doc_ir.blocks:
            if block.block_type in (BlockType.HEADING1, BlockType.HEADING2):
                # Start a new slide for each heading
                slide = prs.slides.add_slide(content_layout)
                slide.shapes.title.text = block.text
                # Prepare content text frame
                if len(slide.placeholders) > 1:
                    current_tf = slide.placeholders[1].text_frame
                    current_tf.clear()
                else:
                    current_tf = None
                current_slide = slide

            elif block.block_type in (BlockType.PARAGRAPH, BlockType.HEADING3):
                if current_tf is not None:
                    p = current_tf.add_paragraph()
                    p.text = block.text
                    p.level = 1 if block.block_type == BlockType.HEADING3 else 0

            elif block.block_type == BlockType.LIST:
                if current_tf is not None:
                    for item in block.items:
                        p = current_tf.add_paragraph()
                        p.text = f"• {item}"
                        p.level = 1

            elif block.block_type == BlockType.TABLE:
                if block.rows and current_slide is not None:
                    rows_count = len(block.rows)
                    cols_count = max(len(r) for r in block.rows)
                    left = Inches(0.5)
                    top = Inches(3.0)
                    width = Inches(9.0)
                    height = Inches(0.4 * rows_count)
                    tbl = current_slide.shapes.add_table(
                        rows_count, cols_count, left, top, width, height
                    ).table
                    for r_idx, row in enumerate(block.rows):
                        for c_idx, cell in enumerate(row):
                            tbl.cell(r_idx, c_idx).text = cell.text

            elif block.block_type == BlockType.PAGE_BREAK:
                # Force a new blank slide
                prs.slides.add_slide(blank_layout)
                current_slide = None
                current_tf = None

        buf = io.BytesIO()
        prs.save(buf)
        logger.debug("pptx_rendered", title=doc_ir.title, slides=len(prs.slides))
        return buf.getvalue()
