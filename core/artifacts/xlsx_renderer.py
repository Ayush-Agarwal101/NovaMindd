"""
NovaMindd — XLSX Renderer

Deterministically renders a DocumentIR into a .xlsx file using openpyxl.
Tables become sheets; paragraphs become text rows.
"""
from __future__ import annotations

import io

from core.artifacts.ir import BlockType, DocumentIR
from core.logging import get_logger

logger = get_logger(__name__)


class XLSXRenderer:
    """Render DocumentIR → bytes (XLSX format)."""

    def render(self, doc_ir: DocumentIR) -> bytes:
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise RuntimeError("openpyxl is required (pip install openpyxl)")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = doc_ir.title[:31]   # sheet name limit

        # Title row
        ws.append([doc_ir.title])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([])

        row_cursor = 3

        for block in doc_ir.blocks:
            if block.block_type in (BlockType.HEADING1, BlockType.HEADING2, BlockType.HEADING3):
                ws.cell(row=row_cursor, column=1, value=block.text).font = Font(bold=True, size=12)
                row_cursor += 1
            elif block.block_type == BlockType.PARAGRAPH:
                ws.cell(row=row_cursor, column=1, value=block.text)
                row_cursor += 1
            elif block.block_type == BlockType.LIST:
                for item in block.items:
                    ws.cell(row=row_cursor, column=1, value=f"• {item}")
                    row_cursor += 1
            elif block.block_type == BlockType.TABLE:
                if block.rows:
                    # Create a new sheet for each table
                    sheet_name = f"Table_{row_cursor}"[:31]
                    tbl_ws = wb.create_sheet(title=sheet_name)
                    header_fill = PatternFill("solid", fgColor="D9E1F2")
                    for r_idx, row in enumerate(block.rows, start=1):
                        for c_idx, cell in enumerate(row, start=1):
                            c = tbl_ws.cell(row=r_idx, column=c_idx, value=cell.text)
                            if cell.bold:
                                c.font = Font(bold=True)
                                c.fill = header_fill
                            # Auto-width (approximate)
                            col_letter = get_column_letter(c_idx)
                            tbl_ws.column_dimensions[col_letter].width = max(
                                tbl_ws.column_dimensions[col_letter].width,
                                min(len(cell.text) + 2, 50),
                            )
                    ws.cell(row=row_cursor, column=1, value=f"[Table → sheet: {sheet_name}]")
                    row_cursor += 1
            elif block.block_type == BlockType.PAGE_BREAK:
                row_cursor += 1

        buf = io.BytesIO()
        wb.save(buf)
        logger.debug("xlsx_rendered", title=doc_ir.title, blocks=len(doc_ir.blocks))
        return buf.getvalue()
